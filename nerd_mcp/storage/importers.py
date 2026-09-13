"""Inventory file decoding and validation independent of SQLite storage."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path

from ..domain.errors import InventoryError
from ..domain.inventory import (
    DEVICE_FIELDS,
    IMPORT_FIELDS,
    TEXT_LIMITS,
    Device,
    valid_host,
    validate_inventory_text,
)


HEADER_ALIASES = {
    "hostname": "device_hostname",
    "devicehostname": "device_hostname",
    "streetaddress": "street_address",
    "zipcode": "zip_code",
    "postal_code": "zip_code",
    "postalcode": "zip_code",
    "serialnumber": "serial_number",
    "type": "device_type",
}


@dataclass(frozen=True)
class InventoryImportRecord:
    row_number: int
    device: Device


@dataclass(frozen=True)
class InventoryImportBatch:
    records: tuple[InventoryImportRecord, ...]
    metadata_fields: frozenset[str]


def _canonical_header(value) -> str:
    text = str(value or "").lstrip("\ufeff").strip().lower()
    text = re.sub(r"[\s-]+", "_", text)
    return HEADER_ALIASES.get(text, HEADER_ALIASES.get(text.replace("_", ""), text))


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_rows(path: str | Path) -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".csv":
        try:
            with source.open(encoding="utf-8-sig", newline="") as stream:
                raw_rows = list(csv.reader(stream, strict=True))
        except (csv.Error, UnicodeError) as exc:
            raise InventoryError(
                "Invalid CSV: save as CSV UTF-8 with properly quoted fields."
            ) from exc
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                raw_rows = [
                    list(row) for row in workbook.active.iter_rows(values_only=True)
                ]
            finally:
                workbook.close()
        except Exception as exc:
            raise InventoryError(
                "Invalid Excel workbook: use an unencrypted .xlsx file."
            ) from exc
    else:
        raise InventoryError(
            "Unsupported inventory file; use a UTF-8 .csv or Excel .xlsx file."
        )
    if not raw_rows:
        raise InventoryError("Inventory file is empty.")
    headers = [_canonical_header(value) for value in raw_rows[0]]
    if not {"name", "host"} <= set(headers):
        raise InventoryError("Inventory header must contain name and host.")
    if "" in headers or len(headers) != len(set(headers)) or set(headers) - IMPORT_FIELDS:
        raise InventoryError(
            "Inventory has duplicate or unknown columns; use the documented headers."
        )
    rows = []
    for number, values in enumerate(raw_rows[1:], start=2):
        if not any(_cell_text(value) for value in values):
            continue
        if len(values) != len(headers):
            raise InventoryError(f"Row {number}: column count does not match header")
        rows.append(
            (number, dict(zip(headers, (_cell_text(value) for value in values))))
        )
    return headers, rows


def load_inventory_import(path: str | Path) -> InventoryImportBatch:
    """Return validated records ready for one repository transaction."""
    headers, rows = _read_rows(path)
    records: list[InventoryImportRecord] = []
    errors: list[str] = []
    seen: set[str] = set()
    for number, row in rows:
        try:
            name, host = row["name"], row["host"]
            profile = (row.get("credential_profile", "") or "default").lower()
            if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,63}", name):
                raise ValueError(
                    "name must be 1-64 letters, digits, dots, underscores or hyphens"
                )
            if name.casefold() in seen:
                raise ValueError("duplicate device name within inventory file")
            seen.add(name.casefold())
            if not valid_host(host):
                raise ValueError("host must be an IP address or DNS hostname")
            try:
                port = int(row.get("port", "") or "22")
            except ValueError:
                raise ValueError("port must be an integer") from None
            if not 1 <= port <= 65535:
                raise ValueError("port must be between 1 and 65535")
            if not re.fullmatch(r"[a-z][a-z0-9_]{0,31}", profile):
                raise ValueError(
                    "credential_profile must start with a letter and use letters, digits or underscores"
                )
            values = {field: "" for field in DEVICE_FIELDS}
            values.update(
                {
                    "name": name,
                    "host": host,
                    "port": port,
                    "credential_profile": profile,
                    "vendor": row.get("vendor", "") or "Cisco",
                    "platform": row.get("platform", "") or "IOS/IOS-XE",
                }
            )
            for field in TEXT_LIMITS:
                if field in row:
                    values[field] = validate_inventory_text(field, row[field])
            country = row.get("country", "").upper()
            if country and not re.fullmatch(r"[A-Z]{3}", country):
                raise ValueError(
                    "country must be a three-letter code such as USA or CAN"
                )
            values["country"] = country
            if values["device_hostname"] and not valid_host(values["device_hostname"]):
                raise ValueError("device_hostname must be a valid hostname")
            if "location" in row and values["location"]:
                values["location_source"] = "csv"
            records.append(InventoryImportRecord(number, Device(**values)))
        except ValueError as exc:
            errors.append(f"Row {number}: {exc}")
    if errors:
        raise InventoryError("\n".join(errors))
    return InventoryImportBatch(
        tuple(records),
        frozenset(headers) - {"name", "host", "port", "credential_profile"},
    )
